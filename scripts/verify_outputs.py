"""Independently validate the default Kaggle input and generated reports."""
from pathlib import Path
import math
import pandas as pd
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
raw = pd.read_csv(root / 'data/raw/online_retail_II.csv', low_memory=False)
d = raw.drop_duplicates()
mask = (~d.Invoice.astype(str).str.upper().str.startswith('C') &
        d.Quantity.gt(0) & d.Price.gt(0) &
        pd.to_datetime(d.InvoiceDate, errors='coerce').notna() & d.Invoice.notna())
expected = (d.loc[mask, 'Quantity'] * d.loc[mask, 'Price']).sum()
summary = pd.read_csv(root / 'reports/summary.csv').set_index('Metric').Value
assert math.isclose(float(summary['Gross sales (GBP)']), expected, abs_tol=.001)
assert int(summary['Sales line items']) == int(mask.sum())
for name in ['monthly_revenue', 'revenue_by_country', 'sales_mix']:
    table = pd.read_csv(root / f'reports/{name}.csv')
    assert math.isclose(table.Revenue.sum(), expected, abs_tol=.001), name
q = pd.read_csv(root / 'reports/data_quality.csv').set_index('Data quality check').Rows
assert q['Rows loaded'] == (q['Exact duplicate rows removed'] +
    q['Rows excluded after deduplication (distinct)'] + q['Valid sales rows analyzed'])
book = load_workbook(root / 'reports/retail_sales_analysis.xlsx')
assert len(book.worksheets) == 9
assert len(book['monthly_revenue']._charts) == 1
assert math.isclose(book['Summary']['B2'].value, expected, abs_tol=.001)
for sheet in book:
    assert sheet.max_row > 1
    assert sheet.freeze_panes == 'A2'
    for row in sheet:
        assert all(c.data_type != 'e' for c in row), sheet.title
print(f'PASS: independent source total £{expected:,.2f}; row reconciliation; aggregate totals; nine workbook sheets and native chart.')

yoy = pd.read_csv(root / 'reports/year_over_year.csv')
assert '2011-12' not in set(yoy.Month)
valid = d.loc[mask].copy()
valid['Month'] = pd.to_datetime(valid.InvoiceDate).dt.strftime('%Y-%m')
valid['Value'] = valid.Quantity * valid.Price
monthly = valid.groupby('Month').Value.sum()
for row in yoy.itertuples():
    prior = str(pd.Period(row.Month, freq='M') - 12)
    assert math.isclose(row.PriorYearRevenue, monthly[prior], abs_tol=.001)
    assert math.isclose(row.RevenueGrowth, monthly[row.Month] / monthly[prior] - 1, abs_tol=1e-9)
print('PASS: all year-over-year values independently recomputed; partial December excluded.')
